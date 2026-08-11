from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import BinaryIO

from openpyxl import load_workbook


class InvoiceFormatError(ValueError):
    pass


MONTHS = {
    "січня": 1, "лютого": 2, "березня": 3, "квітня": 4,
    "травня": 5, "червня": 6, "липня": 7, "серпня": 8,
    "вересня": 9, "жовтня": 10, "листопада": 11, "грудня": 12,
}

HEADER_ALIASES = {
    "code": {"код", "код товару", "артикул"},
    "name": {"товари (послуги)", "товар", "найменування", "назва"},
    "qty": {"кількість", "к-сть"},
    "rate": {"ціна", "ціна за од."},
    "amount": {"сума"},
}


def _text(value) -> str:
    return " ".join(str(value or "").strip().split())


def _decimal(value, label: str, row_no: int) -> Decimal:
    try:
        return Decimal(str(value).replace(" ", "").replace(",", "."))
    except (InvalidOperation, AttributeError):
        raise InvoiceFormatError(f"Рядок {row_no}: некоректне поле «{label}»: {value!r}")


def _document_details(sheet) -> tuple[str | None, date | None]:
    title = " ".join(
        _text(sheet.cell(row=row, column=col).value)
        for row in range(1, min(sheet.max_row, 12) + 1)
        for col in range(1, min(sheet.max_column, 15) + 1)
    )
    number_match = re.search(r"накладн\w*\s*№\s*([^\s]+)", title, re.IGNORECASE)
    date_match = re.search(
        r"(?:від\s*)?(\d{1,2})\s+(" + "|".join(MONTHS) + r")\s+(\d{4})",
        title,
        re.IGNORECASE,
    )
    document_date = None
    if date_match:
        document_date = date(
            int(date_match.group(3)), MONTHS[date_match.group(2).lower()], int(date_match.group(1))
        )
    return (number_match.group(1) if number_match else None, document_date)


def _find_headers(sheet) -> tuple[int, dict[str, int]]:
    for row in range(1, min(sheet.max_row, 80) + 1):
        found: dict[str, int] = {}
        for col in range(1, sheet.max_column + 1):
            value = _text(sheet.cell(row=row, column=col).value).casefold()
            for key, aliases in HEADER_ALIASES.items():
                if value in aliases:
                    found[key] = col
        if set(found) == set(HEADER_ALIASES):
            return row, found
    raise InvoiceFormatError("Не знайдено таблицю товарів із колонками Код, Назва, Кількість, Ціна та Сума.")


def _uom(sheet, row: int, qty_col: int, rate_col: int) -> str:
    for col in range(qty_col + 1, rate_col):
        value = _text(sheet.cell(row=row, column=col).value)
        if value and not re.fullmatch(r"[\d\s.,-]+", value):
            return value
    return "шт"


def parse_mobitime(source: bytes | BinaryIO) -> dict:
    stream = BytesIO(source) if isinstance(source, bytes) else source
    workbook = load_workbook(stream, read_only=False, data_only=True)
    sheet = workbook.active
    header_row, columns = _find_headers(sheet)
    document_no, document_date = _document_details(sheet)
    items = []
    blank_rows = 0

    for row in range(header_row + 1, sheet.max_row + 1):
        code = _text(sheet.cell(row=row, column=columns["code"]).value)
        name = _text(sheet.cell(row=row, column=columns["name"]).value)
        if not code and not name:
            blank_rows += 1
            if items and blank_rows >= 3:
                break
            continue
        blank_rows = 0
        if name.casefold().startswith(("разом", "всього")):
            break
        if not code:
            raise InvoiceFormatError(f"Рядок {row}: у товару відсутній код.")
        qty = _decimal(sheet.cell(row=row, column=columns["qty"]).value, "Кількість", row)
        rate = _decimal(sheet.cell(row=row, column=columns["rate"]).value, "Ціна", row)
        amount = _decimal(sheet.cell(row=row, column=columns["amount"]).value, "Сума", row)
        items.append({
            "row_no": row,
            "supplier_code": code,
            "barcode": code if code.isdigit() and 8 <= len(code) <= 14 else None,
            "item_name": name or code,
            "qty": qty,
            "source_uom": _uom(sheet, row, columns["qty"], columns["rate"]),
            "source_rate": rate,
            "source_amount": amount,
        })

    if not items:
        raise InvoiceFormatError("У таблиці не знайдено жодного товару.")
    return {
        "sheet_name": sheet.title,
        "document_no": document_no,
        "document_date": document_date,
        "items": items,
        "total_qty": sum((row["qty"] for row in items), Decimal("0")),
        "total_amount": sum((row["source_amount"] for row in items), Decimal("0")),
    }
