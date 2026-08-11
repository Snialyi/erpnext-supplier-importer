import unittest
from datetime import date
from io import BytesIO

from openpyxl import Workbook

from supplier_invoice_importer.parsers.mobitime import InvoiceFormatError, parse_mobitime


def workbook_bytes(with_items=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TDSheet"
    sheet.cell(2, 1, "Видаткова накладна № 27/08 від 6 серпня 2026 р.")
    for column, value in {4: "Код", 10: "Товари (Послуги)", 37: "Кількість", 41: "Ціна", 42: "Сума"}.items():
        sheet.cell(14, column, value)
    if with_items:
        sheet.cell(16, 4, "2000000325279")
        sheet.cell(16, 10, "Тестовий товар")
        sheet.cell(16, 37, 3)
        sheet.cell(16, 39, "шт")
        sheet.cell(16, 41, 53.43)
        sheet.cell(16, 42, 160.29)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class MobiTimeParserTestCase(unittest.TestCase):
    def test_parses_document_and_items(self):
        result = parse_mobitime(workbook_bytes())
        self.assertEqual(result["document_no"], "27/08")
        self.assertEqual(result["document_date"], date(2026, 8, 6))
        self.assertEqual(result["items"][0]["barcode"], "2000000325279")
        self.assertEqual(str(result["total_qty"]), "3")
        self.assertEqual(str(result["total_amount"]), "160.29")

    def test_rejects_empty_invoice(self):
        with self.assertRaises(InvoiceFormatError):
            parse_mobitime(workbook_bytes(with_items=False))


if __name__ == "__main__":
    unittest.main()
